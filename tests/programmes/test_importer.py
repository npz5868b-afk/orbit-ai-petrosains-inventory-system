"""B1-only tests. Synthetic fixtures are not official or operational evidence.

Set ORBIT_PROGRAMME_WORKBOOK to run the original-source integration checks.
Run with unittest discovery here to avoid application/DB pytest fixtures.
"""
import copy
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest
import warnings
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location("programme_importer", ROOT / "backend/scripts/import_programme_catalogue.py")
imp = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(imp)


def fixture(path, *, edit=None):
    """A deliberately synthetic one-offering workbook; no live stock or bookings."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Participant_README")
    ws["A1"] = "SYNTHETIC TEST DATA ONLY"
    ws = wb.create_sheet("Data_Dictionary")
    ws.append(imp.DICTIONARY_COLUMNS)
    for field in ("Unknown", "N/A", "To be validated"):
        ws.append([field, "Synthetic definition", "Text", field])
    ws = wb.create_sheet("Offerings_Master")
    ws.append(imp.MASTER_COLUMNS)
    row = {k: "To be validated" for k in imp.MASTER_COLUMNS}
    row.update(Offering_ID="ACT-001", Activity_Title="Synthetic Exact Title!", Internet_Required="Optional", Electricity_Required="No", Min_Participants=0)
    ws.append([row[k] for k in imp.MASTER_COLUMNS])
    ws = wb.create_sheet("Theme_Objective_Mapping")
    ws.append(imp.MAPPING_COLUMNS)
    ws.append(["MAP-001", "Synthetic theme", "concept", "objective", "audience", "ACT-001", "N/A", "reason", "To be validated"])
    ws = wb.create_sheet("Constraint_Rules")
    ws.append(imp.RULE_COLUMNS)
    for i in range(1, 13):
        ws.append([f"RULE-{i:03}", "synthetic", "trigger", "action", "severity", "guidance"])
    ws = wb.create_sheet("ACT-001_Test")
    for cell, val in {"B3": "No", "C3": "Item", "E3": "Number per pack", "F3": "Number of ready pack needed", "G3": "Quantity needed", "H3": "Remark/size/spec", "B4": 1, "C4": "Test cable", "E4": "100 units", "F4": 80, "G4": "1 pack", "H4": "synthetic spec", "B5": 2, "D5": "Displaced material", "E5": 1, "G5": "1 unit for 1 session", "H5": "test spec", "C6": "Unresolved row"}.items():
        ws[cell] = val
    ws.merge_cells("C3:D3")
    ws.merge_cells("F4:F5")
    if edit:
        edit(wb)
    wb.save(path)
    wb.close()


class SyntheticImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "synthetic.xlsx"
        fixture(self.path)

    def build(self):
        return imp.build_dataset(self.path, expected_offerings=1)

    def test_lexical_special_values_remain_distinct(self):
        samples = [(None, "blank"), ("", "empty_text"), (" ", "whitespace_text"), ("Unknown", "unknown"), ("N/A", "not_applicable"), ("Optional", "optional"), ("No", "no"), (0, "zero"), (False, "boolean"), ("To be validated", "to_be_validated"), ("TBC", "undefined_placeholder"), ("To be confirmed", "undefined_placeholder"), ("0", "text"), ("Unknown ", "text")]
        for raw, state in samples:
            with self.subTest(raw=raw):
                self.assertEqual(imp.semantic_state(raw), state)
        data, _ = self.build()
        fields = data["programme_offerings.json"]["records"][0]["fields"]
        self.assertEqual(fields["Internet_Required"]["value"], "Optional")
        self.assertEqual(fields["Electricity_Required"]["value"], "No")
        self.assertEqual(fields["Min_Participants"]["state"], "zero")

    def test_count_mismatch_is_rejected_without_padding(self):
        with self.assertRaisesRegex(imp.ImportValidationError, "Expected 21 offerings, found 1"):
            imp.build_dataset(self.path)

    def test_duplicate_offering_id_rejected(self):
        fixture(self.path, edit=lambda w: w["Offerings_Master"].append([c.value for c in w["Offerings_Master"][2]]))
        with self.assertRaisesRegex(imp.ImportValidationError, "Duplicate Offering_ID"):
            self.build()

    def test_official_name_preserved_and_missing_name_rejected(self):
        data, _ = self.build()
        self.assertEqual(imp.value(data["programme_offerings.json"]["records"][0], "Activity_Title"), "Synthetic Exact Title!")
        fixture(self.path, edit=lambda w: setattr(w["Offerings_Master"]["B2"], "value", None))
        with self.assertRaisesRegex(imp.ImportValidationError, "Missing official Activity_Title"):
            self.build()

    def test_orphan_act_rejected(self):
        fixture(self.path, edit=lambda w: setattr(w["ACT-001_Test"], "title", "ACT-999_Test"))
        with self.assertRaisesRegex(imp.ImportValidationError, "Orphan ACT"):
            self.build()

    def test_missing_act_rejected(self):
        fixture(self.path, edit=lambda w: w.remove(w["ACT-001_Test"]))
        with self.assertRaisesRegex(imp.ImportValidationError, "exactly one ACT"):
            self.build()

    def test_unknown_mapping_reference_rejected(self):
        fixture(self.path, edit=lambda w: setattr(w["Theme_Objective_Mapping"]["F2"], "value", "ACT-999"))
        with self.assertRaisesRegex(imp.ImportValidationError, "Invalid mapping reference"):
            self.build()

    def test_missing_or_incomplete_rule_rejected(self):
        fixture(self.path, edit=lambda w: w["Constraint_Rules"].delete_rows(13))
        with self.assertRaisesRegex(imp.ImportValidationError, "Expected source rules"):
            self.build()
        fixture(self.path, edit=lambda w: setattr(w["Constraint_Rules"]["F2"], "value", None))
        with self.assertRaisesRegex(imp.ImportValidationError, "Incomplete constraint"):
            self.build()

    def test_missing_dictionary_prevents_interpretation(self):
        fixture(self.path, edit=lambda w: w.remove(w["Data_Dictionary"]))
        with self.assertRaisesRegex(imp.ImportValidationError, "Missing required sheets"):
            self.build()

    def test_unexpected_sheet_and_schema_drift_rejected(self):
        fixture(self.path, edit=lambda w: w.create_sheet("Unreviewed_New_Data"))
        with self.assertRaisesRegex(imp.ImportValidationError, "Unexpected sheet"):
            self.build()
        fixture(self.path, edit=lambda w: setattr(w["Offerings_Master"]["A1"], "value", "Renamed_ID"))
        with self.assertRaisesRegex(imp.ImportValidationError, "missing columns"):
            self.build()

    def test_pack_quantity_and_session_context_are_not_scaled(self):
        data, _ = self.build()
        rows = data["programme_materials.json"]["sheets"][0]["records"]
        self.assertEqual(rows[0]["fields"]["Number per pack"][0]["value"], "100 units")
        self.assertEqual(rows[0]["fields"]["Quantity needed"][0]["value"], "1 pack")
        self.assertEqual(rows[1]["fields"]["Quantity needed"][0]["value"], "1 unit for 1 session")
        for row in rows:
            for key in ("operational_usage_basis", "inventory_mapping", "consumption_or_reuse", "pack_conversion"):
                self.assertIsNone(row[key])

    def test_displaced_name_and_merged_blank_are_traceable(self):
        data, _ = self.build()
        row = data["programme_materials.json"]["sheets"][0]["records"][1]
        self.assertEqual(row["item_name"], "Displaced material")
        self.assertEqual(row["fields"]["Item"][0]["cell"], "D5")
        merged = row["fields"]["Number of ready pack needed"][0]
        self.assertIsNone(merged["value"])
        self.assertEqual(merged["state"], "blank")
        self.assertEqual(merged["merged_context"], {"range": "F4:F5", "anchor": "F4", "anchor_value": 80})

    def test_ambiguous_item_rows_are_retained(self):
        data, _ = self.build()
        row = data["programme_materials.json"]["sheets"][0]["records"][2]
        self.assertEqual(row["item_name"], "Unresolved row")
        self.assertEqual(row["row_kind"], "material_or_section_needs_review")

    def test_shifted_material_header_detected(self):
        fixture(self.path, edit=lambda w: w["ACT-001_Test"].insert_rows(1, 16))
        data, _ = self.build()
        sheet = data["programme_materials.json"]["sheets"][0]
        self.assertEqual(sheet["header_row"], 19)
        self.assertEqual(sheet["records"][0]["source"]["row"], 20)

    def test_missing_formula_cache_is_not_calculated(self):
        fixture(self.path, edit=lambda w: setattr(w["ACT-001_Test"]["G4"], "value", "=80*100"))
        data, _ = self.build()
        c = data["programme_materials.json"]["sheets"][0]["records"][0]["fields"]["Quantity needed"][0]
        self.assertEqual(c["formula"], "=80*100")
        self.assertIsNone(c["value"])
        self.assertEqual(c["state"], "formula_cache_missing")
        self.assertFalse(c["recalculated"])
        self.assertEqual(data["import_audit.json"]["summary"]["formula_caches_missing"], 1)

    def test_stored_cache_not_recomputed_even_when_different(self):
        fixture(self.path, edit=lambda w: setattr(w["ACT-001_Test"]["G4"], "value", "=80*100"))
        # Set a deliberately stale stored cache using OOXML in the synthetic fixture.
        with ZipFile(self.path) as archive:
            contents = {n: archive.read(n) for n in archive.namelist()}
        member = "xl/worksheets/sheet6.xml"
        root = ET.fromstring(contents[member])
        cell = root.find(".//s:c[@r='G4']", imp.NS)
        cell.find("s:v", imp.NS).text = "7"
        contents[member] = ET.tostring(root)
        with ZipFile(self.path, "w", ZIP_DEFLATED) as archive:
            for name, payload in contents.items():
                archive.writestr(name, payload)
        data, _ = self.build()
        c = data["programme_materials.json"]["sheets"][0]["records"][0]["fields"]["Quantity needed"][0]
        self.assertEqual(c["value"], 7)
        self.assertEqual(c["formula"], "=80*100")
        self.assertEqual(c["cache_status"], "present")
        self.assertFalse(c["recalculated"])

    def test_excel_error_is_not_zero_or_missing(self):
        fixture(self.path, edit=lambda w: setattr(w["ACT-001_Test"]["G4"], "value", "#VALUE!"))
        data, _ = self.build()
        c = data["programme_materials.json"]["sheets"][0]["records"][0]["fields"]["Quantity needed"][0]
        self.assertEqual((c["state"], c["value"]), ("excel_error", "#VALUE!"))

    def test_text_in_price_column_is_retained_and_flagged(self):
        def edit(w):
            w["ACT-001_Test"]["I3"] = "Price/unit"
            w["ACT-001_Test"]["I4"] = "Tall Glass"
        fixture(self.path, edit=edit)
        data, _ = self.build()
        c = data["programme_materials.json"]["sheets"][0]["records"][0]["fields"]["Price/unit"][0]
        self.assertEqual(c["value"], "Tall Glass")
        self.assertTrue(any(i["code"] == "nonnumeric_price_cell" and i["source"]["cell"] == "I4" for i in data["import_audit.json"]["issues"]))

    def test_deterministic_and_source_not_modified(self):
        before = self.path.read_bytes()
        a = self.build()
        b = self.build()
        self.assertEqual(imp.canonical_bytes(a), imp.canonical_bytes(b))
        self.assertEqual(self.path.read_bytes(), before)


@unittest.skipUnless(os.environ.get("ORBIT_PROGRAMME_WORKBOOK"), "Set ORBIT_PROGRAMME_WORKBOOK for official-source integration tests")
class OfficialWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.path = Path(os.environ["ORBIT_PROGRAMME_WORKBOOK"])
        cls.before = hashlib.sha256(cls.path.read_bytes()).hexdigest()
        cls.data, cls.manifest = imp.build_dataset(cls.path)

    def test_official_source_hash_counts_and_names(self):
        self.assertEqual(self.before, "10810eb3e4ed3fa3a2cb94c828fa7965fbe4a03370157bc5b4a7ab9a705b5d4e")
        self.assertEqual(len(self.manifest["sheet_names"]), 26)
        rows = self.data["programme_offerings.json"]["records"]
        self.assertEqual({imp.value(r, "Offering_ID") for r in rows}, {f"ACT-{i:03}" for i in range(1, 22)})
        self.assertEqual(imp.value(rows[19], "Activity_Title"), "Fabolous Fizzy")
        summary = self.data["import_audit.json"]["summary"]
        for key, expected in {"offering_count": 21, "constraint_count": 12, "theme_mapping_count": 25, "act_sheet_count": 21, "material_candidate_rows": 275, "material_rows_with_quantity": 254, "ambiguous_material_or_section_rows": 21, "formula_count": 17, "formula_caches_missing": 0, "excel_error_count": 34}.items():
            self.assertEqual(summary[key], expected, key)

    def test_each_table_data_cell_matches_original(self):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            wb = openpyxl.load_workbook(self.path, read_only=True, data_only=False)
            try:
                for sheet, output in [("Offerings_Master", "programme_offerings.json"), ("Constraint_Rules", "programme_constraints.json"), ("Theme_Objective_Mapping", "programme_theme_mapping.json"), ("Data_Dictionary", "programme_dictionary.json")]:
                    rows = list(wb[sheet].iter_rows())
                    expected = {c.coordinate: c.value for row in rows[1:] for c in row if c.value is not None}
                    actual = {c["cell"]: c["raw_value"] for r in self.data[output]["records"] for c in r["fields"].values() if c["raw_value"] is not None}
                    self.assertEqual(actual, expected, sheet)
            finally:
                wb.close()

    def test_every_act_nonempty_cell_is_preserved_with_coordinate(self):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            wb = openpyxl.load_workbook(self.path, read_only=True, data_only=False)
            try:
                for sheet in self.data["programme_materials.json"]["sheets"]:
                    atoms = list(sheet["headers"])
                    for row in sheet["preamble_rows"] + sheet["other_rows"]:
                        atoms.extend(row["cells"])
                    for row in sheet["records"]:
                        self.assertEqual(row["source"]["sheet"], sheet["sheet"])
                        self.assertEqual(row["Offering_ID"], sheet["Offering_ID"])
                        for cells in row["fields"].values():
                            atoms.extend(cells)
                    actual = {c["cell"]: c["raw_value"] for c in atoms if c["raw_value"] is not None}
                    expected = {c.coordinate: c.value for row in wb[sheet["sheet"]] for c in row if c.value is not None}
                    self.assertEqual(actual, expected, sheet["sheet"])
            finally:
                wb.close()

    def test_all_twelve_constraint_rules_preserved(self):
        rules = self.data["programme_constraints.json"]["records"]
        self.assertEqual({imp.value(r, "Rule_ID") for r in rules}, {f"RULE-{i:03}" for i in range(1, 13)})
        self.assertTrue(all(all(c["raw_value"] is not None for c in r["fields"].values()) for r in rules))

    def test_original_formula_caches_are_traced_without_recalculation(self):
        fs = self.data["import_audit.json"]["formulas"]
        self.assertEqual(len(fs), 17)
        self.assertTrue(all(f["cache_status"] == "present" and f["recalculated"] is False for f in fs))
        f = next(f for f in fs if f["source"]["sheet"] == "ACT-006_LED Lantern")
        self.assertEqual((f["cell"], f["formula"], f["value"]), ("G12", "=F12*E12", 80))

    def test_original_usage_fields_remain_unverified(self):
        for sheet in self.data["programme_materials.json"]["sheets"]:
            for row in sheet["records"]:
                for key in ("operational_usage_basis", "inventory_mapping", "consumption_or_reuse", "pack_conversion"):
                    self.assertIsNone(row[key])
        self.assertEqual(self.data["import_audit.json"]["operational_validation"], "needs_verification")
        self.assertIsNone(self.manifest["official_release_or_update_date"])

    def test_exports_are_byte_identical_and_manifest_hashes_match(self):
        acquisition = json.loads((ROOT / "backend/data/programmes/source_manifest.json").read_text(encoding="utf-8"))["acquisition"]
        with tempfile.TemporaryDirectory() as tmp:
            a, b = Path(tmp) / "a", Path(tmp) / "b"
            imp.export_dataset(self.path, a, acquisition)
            imp.export_dataset(self.path, b, acquisition)
            for name in sorted(p.name for p in a.iterdir()):
                self.assertEqual((a / name).read_bytes(), (b / name).read_bytes(), name)
                self.assertEqual((a / name).read_bytes(), (ROOT / "backend/data/programmes" / name).read_bytes(), name)
            manifest = json.loads((a / "source_manifest.json").read_text(encoding="utf-8"))
            for name, digest in manifest["outputs"].items():
                self.assertEqual(hashlib.sha256((a / name).read_bytes()).hexdigest(), digest)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), self.before)

    def test_bad_acquisition_and_changed_outputs_not_overwritten(self):
        acquisition = json.loads((ROOT / "backend/data/programmes/source_manifest.json").read_text(encoding="utf-8"))["acquisition"]
        with tempfile.TemporaryDirectory() as tmp:
            bad = copy.deepcopy(acquisition)
            bad["sha256"] = "0" * 64
            with self.assertRaisesRegex(imp.ImportValidationError, "SHA-256"):
                imp.export_dataset(self.path, Path(tmp) / "bad", bad)
            self.assertFalse((Path(tmp) / "bad").exists())
            out = Path(tmp) / "out"
            out.mkdir()
            target = out / "programme_offerings.json"
            target.write_text("existing work", encoding="utf-8")
            with self.assertRaisesRegex(imp.ImportValidationError, "Refusing to overwrite"):
                imp.export_dataset(self.path, out, acquisition)
            self.assertEqual(target.read_text(encoding="utf-8"), "existing work")
            self.assertEqual(len(list(out.iterdir())), 1)


if __name__ == "__main__":
    unittest.main()
