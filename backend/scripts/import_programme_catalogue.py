"""Source-addressed B1 cell-data import. No recommendations, DB access or recalculation.

Requires openpyxl==3.1.5. See docs/programmes/data-contract.md for JSON semantics.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import posixpath
import re
import sys
import warnings
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries

SCHEMA_VERSION = "1.0.0"
IMPORTER_VERSION = "1.0.0"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REQUIRED_SHEETS = ("Participant_README", "Data_Dictionary", "Offerings_Master",
                   "Theme_Objective_Mapping", "Constraint_Rules")
MASTER_COLUMNS = "Offering_ID Activity_Title Offering_Type Short_Description STEM_Domain Key_Concepts Learning_Outcomes Recommended_Age Audience_Types Standard_Duration_Min Min_Participants Max_Participants Delivery_Mode Indoor_Outdoor Electricity_Required Internet_Required Water_Required Facilitators_Required Setup_Time_Min Accessibility_Notes Safety_Level Key_Hazards Participant_Handling_Rule Suitable_Themes Suitable_Objectives Engagement_Methods Customisable_Elements Key_Constraints Complementary_Activities Cost_Band Availability_Status Data_Confidence Last_Validated Notes".split()
RULE_COLUMNS = "Rule_ID Rule_Type Trigger Recommended_Action Severity Guidance".split()
MAPPING_COLUMNS = "Mapping_ID Theme Related_Concepts Stakeholder_Objective Suitable_Audience Primary_Offering_ID Secondary_Offering_ID Reason_for_Match Validation_Status".split()
DICTIONARY_COLUMNS = ["Field_Name", "Definition", "Allowed_Values_or_Format", "Example"]
QUANTITY_HEADERS = {"Number per pack", "Number of ready pack needed", "Quantity needed",
                    "Quantity (unit) in 1 set", "Quantity for 40 sets", "Quantity per Set",
                    "Quantity", "Total unit / session"}
SPEC_HEADERS = {"Remark/size/spec", "Specification", "Suggested Specification"}


class ImportValidationError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise ImportValidationError(message)


def scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    require(value is None or isinstance(value, (str, int, float, bool)),
            f"Unsupported Excel scalar type: {type(value).__name__}")
    return value


def semantic_state(value, cell_type=None):
    # These are lexical tags, NOT conversions or feasibility/approval decisions.
    if cell_type == "e":
        return "excel_error"
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "zero" if value == 0 else "number"
    if value == "":
        return "empty_text"
    if not value.strip():
        return "whitespace_text"
    exact = {"Unknown": "unknown", "N/A": "not_applicable",
             "To be validated": "to_be_validated", "Optional": "optional",
             "No": "no", "Yes": "yes"}
    if value in exact:
        return exact[value]
    if value.strip().lower() in {"tbc", "to be confirmed", "nil"}:
        return "undefined_placeholder"
    return "text"


def atom(cell, cached_cell=None):
    raw = scalar(cell.value)
    if cell.data_type == "f":
        cached = scalar(cached_cell.value) if cached_cell else None
        cache_type = cached_cell.data_type if cached_cell else None
        state = "formula_cache_missing" if cached is None else semantic_state(cached, cache_type)
        return {"cell": cell.coordinate, "raw_value": raw, "value": cached,
                "state": state, "cell_type": "f", "formula": raw,
                "cache_status": "missing" if cached is None else "error" if cache_type == "e" else "present",
                "recalculated": False}
    return {"cell": cell.coordinate, "raw_value": raw, "value": raw,
            "state": semantic_state(raw, cell.data_type), "cell_type": cell.data_type}


def blank(coordinate):
    return {"cell": coordinate, "raw_value": None, "value": None,
            "state": "blank", "cell_type": "n"}


def xlsx_layout(path):
    """Read merges without loading drawings, modifying or re-saving the workbook."""
    with ZipFile(path) as z:
        rel = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        targets = {r.attrib["Id"]: r.attrib["Target"] for r in rel}
        book = ET.fromstring(z.read("xl/workbook.xml"))
        layouts = {}
        for s in book.findall("s:sheets/s:sheet", NS):
            rid = s.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
            target = targets[rid]
            member = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            tree = ET.fromstring(z.read(member))
            layouts[s.attrib["name"]] = {"xml_member": member,
                "merged_ranges": [m.attrib["ref"] for m in tree.findall("s:mergeCells/s:mergeCell", NS)]}
        return layouts


def read_source(path):
    layouts = xlsx_layout(path)
    with warnings.catch_warnings(record=True) as captured:
        warnings.simplefilter("always")
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            require(all(n in formulas.sheetnames for n in REQUIRED_SHEETS),
                    f"Missing required sheets: {set(REQUIRED_SHEETS) - set(formulas.sheetnames)}")
            sheets = {}
            # Guidance is deliberately consumed before activity fields.
            order = ["Participant_README", "Data_Dictionary"] + [n for n in formulas.sheetnames if n not in {"Participant_README", "Data_Dictionary"}]
            for name in order:
                rows = {}
                for number, (r, cr) in enumerate(zip(formulas[name], cached[name]), 1):
                    cells = {c.coordinate: atom(c, cc) for c, cc in zip(r, cr) if c.value is not None}
                    if cells:
                        rows[number] = cells
                sheets[name] = {**layouts[name], "rows": rows}
            return sheets, list(formulas.sheetnames), sorted({str(w.message) for w in captured})
        finally:
            formulas.close()
            cached.close()


def value(record, name):
    return record["fields"][name]["value"]


def table(sheets, name, columns):
    source = sheets[name]
    header = source["rows"].get(1, {})
    by_name = {}
    for c in header.values():
        label = c["value"]
        require(isinstance(label, str) and label not in by_name, f"{name}: duplicate/invalid header {label!r}")
        by_name[label] = re.sub(r"\d+$", "", c["cell"])
    require(set(columns) <= set(by_name), f"{name}: missing columns {set(columns) - set(by_name)}")
    records = []
    for rn, cells in source["rows"].items():
        if rn == 1:
            continue
        known_cells = {f"{col}{rn}" for col in by_name.values()}
        require(set(cells) <= known_cells, f"{name}!{rn}: data under an unlabelled column")
        records.append({"source": {"sheet": name, "row": rn}, "fields": {
            key: cells.get(f"{col}{rn}", blank(f"{col}{rn}")) for key, col in by_name.items()}})
    return {"columns": list(by_name), "records": records}


def unique(records, key, pattern):
    ids = [value(r, key) for r in records]
    require(all(isinstance(v, str) and re.fullmatch(pattern, v) for v in ids), f"Invalid or missing {key}")
    require(len(ids) == len(set(ids)), f"Duplicate {key}")
    return set(ids)


def material_sheet(name, sheet, ids):
    offering_id = name[:7]
    require(offering_id in ids, f"Orphan ACT sheet: {name}")
    headers = [(rn, cs) for rn, cs in sheet["rows"].items()
               if any(c["value"] in ("Item", "Item Name") for c in cs.values())
               and any(c["value"] in ("No", "No.") for c in cs.values())]
    require(len(headers) == 1, f"{name}: expected one identifiable material header, found {len(headers)}")
    hr, hc = headers[0]
    labels = []
    for c in hc.values():
        col = range_boundaries(c["cell"])[0]
        labels.append((col, c["value"]))
    labels.sort()
    require(len({label for _, label in labels}) == len(labels), f"{name}: duplicate material header")
    require(any(label in QUANTITY_HEADERS for _, label in labels), f"{name}: unrecognised quantity layout")
    merge_ranges = [(r, range_boundaries(r)) for r in sheet["merged_ranges"]]
    records, other_rows = [], []
    section = None
    for rn, cells in sheet["rows"].items():
        if rn <= hr:
            continue
        fields = {}
        for idx, (col, label) in enumerate(labels):
            end = labels[idx + 1][0] if idx + 1 < len(labels) else max(col + 1, max(range_boundaries(k)[0] for k in cells) + 1)
            atoms = [cells[f"{get_column_letter(c)}{rn}"] for c in range(col, end) if f"{get_column_letter(c)}{rn}" in cells]
            if not atoms:
                item = blank(f"{get_column_letter(col)}{rn}")
                for merged, (left, top, right, bottom) in merge_ranges:
                    if left <= col <= right and top <= rn <= bottom:
                        anchor = f"{get_column_letter(left)}{top}"
                        item["merged_context"] = {"range": merged, "anchor": anchor,
                            "anchor_value": sheet["rows"].get(top, {}).get(anchor, {}).get("value")}
                        break
                atoms = [item]
            fields[label] = atoms
        names = fields.get("Item", fields.get("Item Name", []))
        actual_names = [c for c in names if c["raw_value"] is not None]
        source = {"sheet": name, "row": rn}
        if not actual_names:
            section = {"source": source, "cells": list(cells.values())}
            other_rows.append({"kind": "section_or_note", **section})
            continue
        has_quantity = any(c["raw_value"] is not None for k, cs in fields.items() if k in QUANTITY_HEADERS for c in cs)
        kind = "material" if has_quantity else "material_or_section_needs_review"
        records.append({"material_row_id": f"{offering_id}:r{rn}", "Offering_ID": offering_id,
            "source": source, "row_kind": kind, "item_name": actual_names[0]["value"] if len(actual_names) == 1 else None,
            "fields": fields, "section_context": section,
            "operational_usage_basis": None, "inventory_mapping": None,
            "consumption_or_reuse": None, "pack_conversion": None,
            "verification_status": "needs_verification"})
        require(set(cells) <= {a["cell"] for aa in fields.values() for a in aa}, f"{name}!{rn}: unconsumed source cell")
    return {"Offering_ID": offering_id, "sheet": name, "header_row": hr,
            "headers": list(hc.values()), "merged_ranges": sheet["merged_ranges"],
            "preamble_rows": [{"row": rn, "cells": list(cs.values())} for rn, cs in sheet["rows"].items() if rn < hr],
            "records": records, "other_rows": other_rows}


def build_dataset(path, *, expected_offerings=21):
    path = Path(path)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    sheets, names, reader_warnings = read_source(path)
    dictionary = table(sheets, "Data_Dictionary", DICTIONARY_COLUMNS)
    definitions = {value(r, "Field_Name"): r for r in dictionary["records"]}
    require(len(definitions) == len(dictionary["records"]), "Duplicate dictionary field")
    require({"Unknown", "N/A", "To be validated"} <= set(definitions), "Missing required dictionary special-value definitions")
    readme = [{"source": {"sheet": "Participant_README", "row": rn}, "cells": list(cells.values())}
              for rn, cells in sheets["Participant_README"]["rows"].items()]
    master = table(sheets, "Offerings_Master", MASTER_COLUMNS)
    ids = unique(master["records"], "Offering_ID", r"ACT-\d{3}")
    require(len(ids) == expected_offerings, f"Expected {expected_offerings} offerings, found {len(ids)}; no records added or removed")
    require(all(isinstance(value(r, "Activity_Title"), str) and value(r, "Activity_Title").strip() for r in master["records"]), "Missing official Activity_Title")
    mapping = table(sheets, "Theme_Objective_Mapping", MAPPING_COLUMNS)
    unique(mapping["records"], "Mapping_ID", r"MAP-\d{3}")
    for r in mapping["records"]:
        for key in ("Primary_Offering_ID", "Secondary_Offering_ID"):
            v = value(r, key)
            require(v in ids or r["fields"][key]["state"] in {"blank", "not_applicable", "unknown", "to_be_validated"}, f"Invalid mapping reference {v!r}")
    rules = table(sheets, "Constraint_Rules", RULE_COLUMNS)
    rule_ids = unique(rules["records"], "Rule_ID", r"RULE-\d{3}")
    require(rule_ids == {f"RULE-{i:03}" for i in range(1, 13)}, "Expected source rules RULE-001 through RULE-012; investigate missing/extra rules")
    require(all(all(value(r, k) is not None and str(value(r, k)).strip() for k in RULE_COLUMNS) for r in rules["records"]), "Incomplete constraint rule")
    acts = [material_sheet(n, sheets[n], ids) for n in names if re.match(r"ACT-\d{3}_", n)]
    require(Counter(a["Offering_ID"] for a in acts) == Counter({i: 1 for i in ids}), "Every offering must have exactly one ACT sheet")
    require(set(names) == set(REQUIRED_SHEETS) | {a["sheet"] for a in acts}, "Unexpected sheet(s); inspect before defining extraction")
    issues = []
    def issue(code, source, detail):
        issues.append({"code": code, "source": source, "detail": detail})
    enums = {k: {v.strip() for v in value(definitions[k], "Allowed_Values_or_Format").split(";")}
             for k in ("Offering_Type", "Delivery_Mode", "Safety_Level", "Cost_Band", "Availability_Status", "Data_Confidence") if k in definitions}
    for r in master["records"]:
        for k, allowed in enums.items():
            if value(r, k) not in allowed:
                issue("dictionary_value_mismatch", {**r["source"], "cell": r["fields"][k]["cell"]}, {"field": k, "actual": value(r, k), "allowed": sorted(allowed)})
        issue("operational_evidence_not_verified", r["source"], {"Offering_ID": value(r, "Offering_ID"), "Availability_Status": value(r, "Availability_Status"), "Notes": value(r, "Notes")})
        for k in ("Recommended_Age", "Standard_Duration_Min", "Min_Participants", "Max_Participants", "Facilitators_Required", "Setup_Time_Min"):
            c = r["fields"][k]
            if c["state"] in {"blank", "unknown", "to_be_validated", "undefined_placeholder", "excel_error", "formula_cache_missing"}:
                issue("unresolved_offering_field", {**r["source"], "cell": c["cell"]}, k)
    if "Activity_Based_Inventory" not in names:
        for row in readme:
            for c in row["cells"]:
                if "Activity_Based_Inventory" in str(c["value"]):
                    issue("readme_sheet_reference_mismatch", {**row["source"], "cell": c["cell"]}, "Activity_Based_Inventory is absent; ACT sheets retained individually")
    formulas, errors, counts = [], [], Counter()
    for n in names:
        for rn, row in sheets[n]["rows"].items():
            for c in row.values():
                source = {"sheet": n, "row": rn, "cell": c["cell"]}
                counts[c["state"]] += 1
                if c["cell_type"] == "f":
                    formulas.append({"source": source, **c})
                    if c["cache_status"] != "present":
                        issue("formula_cache_unavailable", source, c)
                if c["state"] == "excel_error":
                    errors.append({"source": source, "value": c["value"]})
                if c["state"] == "undefined_placeholder":
                    issue("undefined_placeholder", source, c["value"])
                if isinstance(c["value"], str) and "\ufffd" in c["value"]:
                    issue("source_text_encoding_damage", source, c["value"])
    for a in acts:
        for r in a["records"]:
            if r["row_kind"] != "material":
                issue("material_or_section_ambiguous", r["source"], r["item_name"])
            if not any(c["raw_value"] is not None for k, cs in r["fields"].items() if k in SPEC_HEADERS for c in cs):
                issue("material_specification_missing", r["source"], r["item_name"])
            for label, cs in r["fields"].items():
                populated = [c for c in cs if c["raw_value"] is not None]
                if len(populated) > 1:
                    issue("multiple_cells_under_header_span", r["source"], {"header": label, "cells": [c["cell"] for c in populated], "interpretation": "Preserved as separate cells; column grouping is not a verified semantic merge"})
                if label == "Price/unit":
                    for c in populated:
                        if not isinstance(c["value"], (int, float)) or isinstance(c["value"], bool):
                            issue("nonnumeric_price_cell", {**r["source"], "cell": c["cell"]}, c["value"])
    summary = {"offering_count": len(ids), "constraint_count": len(rules["records"]),
        "theme_mapping_count": len(mapping["records"]), "act_sheet_count": len(acts),
        "material_candidate_rows": sum(len(a["records"]) for a in acts),
        "material_rows_with_quantity": sum(r["row_kind"] == "material" for a in acts for r in a["records"]),
        "ambiguous_material_or_section_rows": sum(r["row_kind"] != "material" for a in acts for r in a["records"]),
        "section_or_note_rows": sum(len(a["other_rows"]) for a in acts),
        "formula_count": len(formulas), "formula_caches_missing": sum(f["cache_status"] == "missing" for f in formulas),
        "formula_cache_errors": sum(f["cache_status"] == "error" for f in formulas),
        "excel_error_count": len(errors), "state_counts_nonblank_source_cells": dict(sorted(counts.items()))}
    common = {"schema_version": SCHEMA_VERSION, "source_sha256": sha}
    data = {
        "programme_offerings.json": {**common, **master},
        "programme_theme_mapping.json": {**common, "use": "supporting_guidance_not_fixed_answers", **mapping},
        "programme_constraints.json": {**common, **rules},
        "programme_dictionary.json": {**common, **dictionary, "participant_readme": readme},
        "programme_materials.json": {**common, "sheets": acts},
        "import_audit.json": {**common, "import_integrity": "pass", "operational_validation": "needs_verification",
            "summary": summary, "formulas": formulas, "excel_errors": errors, "issues": issues,
            "reader_warnings": reader_warnings,
            "limitations": ["No formula engine executed; cached values may be stale", "Images/unsupported extensions not rendered or exported; original workbook never re-saved", "All material usage basis, reuse, inventory mapping and pack conversion remain unverified", "No official programme price, stock, future bookings, personnel, venue or safety approval created"]},
    }
    manifest = {**common, "importer_version": IMPORTER_VERSION, "workbook_filename": path.name,
        "workbook_size_bytes": path.stat().st_size, "sheet_names": names,
        "sheet_layouts": [{"sheet": n, "nonempty_rows": len(sheets[n]["rows"]),
            "nonempty_cells": sum(len(r) for r in sheets[n]["rows"].values()), **{k: v for k, v in sheets[n].items() if k != "rows"}} for n in names],
        "expected_offering_count": expected_offerings, "formula_mode": "formula_and_stored_cache_read_separately_no_recalculation",
        "official_release_or_update_date": None}
    return data, manifest


def canonical_bytes(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False) + "\n").encode("utf-8")


def export_dataset(path, output, acquisition):
    require(set(acquisition) == {"source_url", "source_file_url", "obtained_at", "source_display_name", "downloaded_filename", "sha256"}, "Acquisition metadata fields do not match contract")
    require(datetime.fromisoformat(acquisition["obtained_at"]).tzinfo is not None, "obtained_at requires timezone")
    data, manifest = build_dataset(path)
    require(acquisition["sha256"] == manifest["source_sha256"], "Acquisition SHA-256 does not match workbook")
    require(acquisition["downloaded_filename"] == Path(path).name, "Acquisition filename does not match workbook")
    manifest["acquisition"] = acquisition
    manifest["outputs"] = {name: hashlib.sha256(canonical_bytes(payload)).hexdigest() for name, payload in sorted(data.items())}
    data["source_manifest.json"] = manifest
    output = Path(output)
    encoded = {name: canonical_bytes(payload) for name, payload in data.items()}
    # Never overwrite a changed file silently; use a new directory for a new source.
    for name, content in encoded.items():
        target = output / name
        require(not target.exists() or target.read_bytes() == content, f"Refusing to overwrite differing output {target}")
    output.mkdir(parents=True, exist_ok=True)
    for name, content in encoded.items():
        (output / name).write_bytes(content)
    return data


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("workbook", type=Path)
    p.add_argument("--output", required=True, type=Path)
    p.add_argument("--acquisition", required=True, type=Path)
    args = p.parse_args()
    try:
        data = export_dataset(args.workbook, args.output, json.loads(args.acquisition.read_text(encoding="utf-8")))
    except (ImportValidationError, OSError, ValueError) as exc:
        print(f"IMPORT FAILED: {exc}", file=sys.stderr)
        return 1
    print(json.dumps(data["import_audit.json"]["summary"], ensure_ascii=False, indent=2))
    print("Import integrity PASS; operational evidence NEEDS VERIFICATION. No DB or original workbook changed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
